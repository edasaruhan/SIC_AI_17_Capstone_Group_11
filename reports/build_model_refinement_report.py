from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
FIGURES = REPORTS / "figures"
DATA = ROOT / "data"
RESULTS = REPORTS / "results"
SUBMISSIONS = ROOT / "submissions"
OUTPUT = REPORTS / "Model Refinement and Test Submission.docx"

DARK_BLUE = "24556A"
PALE_BLUE = "EAF1F4"
PALE_GRAY = "F5F5F5"
BORDER = "D9D9D9"
BLACK = "000000"
WHITE = "FFFFFF"


def load_json(path: Path):
    return json.loads(path.read_text(encoding="utf-8"))


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), BORDER)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def set_font(run, name="Aptos", size=None, bold=None, italic=None, color=BLACK) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_font(run, size=8, color="666666")
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
    title.font.size = Pt(23)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(9)
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    for style_name, size, before, after in (
        ("Heading 1", 16, 14, 6),
        ("Heading 2", 12.5, 10, 4),
        ("Heading 3", 11, 8, 3),
    ):
        style = styles[style_name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Aptos Display")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.text = "SIC AI 17 Capstone Group 11"
    footer_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in footer_p.runs:
        set_font(run, size=8, color="666666")
    add_page_number(footer.add_paragraph())


def add_paragraph(doc: Document, text: str, *, bold_lead: str | None = None):
    paragraph = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_font(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead):])
        set_font(rest)
    else:
        run = paragraph.add_run(text)
        set_font(run)
    return paragraph


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        paragraph = doc.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(item)
        set_font(run)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, value in enumerate(headers):
        cell = header.cells[index]
        cell.text = value
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, DARK_BLUE)
        set_cell_margins(cell)
        if widths:
            cell.width = widths[index]
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                set_font(run, size=9.2, bold=True, color=WHITE)
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for col_index, value in enumerate(values):
            cell = row.cells[col_index]
            cell.text = str(value)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = widths[col_index]
            if row_index % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if col_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in paragraph.runs:
                    set_font(run, size=9.1)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_figure(doc: Document, filename: str, caption: str, width=6.35, alt_text=None) -> None:
    path = FIGURES / filename
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    inline = run.add_picture(str(path), width=Inches(width))
    doc_pr = inline._inline.docPr
    doc_pr.set("descr", alt_text or caption)
    caption_p = doc.add_paragraph()
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.space_after = Pt(7)
    caption_run = caption_p.add_run(caption)
    set_font(caption_run, size=8.8, italic=True, color="555555")


def add_code(doc: Document, lines: list[str]) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.25)
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(6)
    for index, line in enumerate(lines):
        run = paragraph.add_run(line + ("\n" if index < len(lines) - 1 else ""))
        set_font(run, name="Consolas", size=8.4)


def add_title_page(doc: Document, walk, submission) -> None:
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("AI for Food Waste Reduction Model Refinement and Test Submission")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_run = subtitle.add_run("Samsung Innovation Campus AI 17 Capstone Group 11")
    set_font(subtitle_run, size=13, bold=True)

    team = doc.add_paragraph()
    team.alignment = WD_ALIGN_PARAGRAPH.CENTER
    team_run = team.add_run("Osman Furkan Erkan  |  Nermin Kılıçarslan  |  Dilara Kuyar  |  Ahsen Nisa Sarıkaya")
    set_font(team_run, size=9.5, color="555555")

    date = doc.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_run = date.add_run("Final submission 17 September 2026")
    set_font(date_run, size=9.5, color="555555")

    doc.add_heading("Executive Summary", level=1)
    add_paragraph(
        doc,
        "Bu rapor, LightGBM talep tahmini modelinin son iyileştirme ve test submission aşamalarını belgeler. Optuna ile ayarlanan model tek holdout döneminde RMSLE değerini 0.47628'den 0.47534'e indirdi. Kalibre P8-P92 quantile bandı %80 hedefi için %80.59 coverage üretti. Üç fold recursive walk-forward doğrulamasında ortalama RMSLE 0.49833 oldu ve persistence baseline'a göre ortalama %38.32 iyileşme sağlandı. Son model haftalar 1-145 üzerindeki 456.548 satırla eğitildi ve hedef içermeyen haftalar 146-155 için 32.573 satırlık submission oluşturdu."
    )
    add_table(
        doc,
        ["Final evidence", "Result", "Artifact"],
        [
            ["Tuned holdout evaluation", "RMSLE 0.47534", "data/hpo_final_comparison.json"],
            ["Calibrated uncertainty band", "Coverage %80.59", "data/quantile_final_calibrated.json"],
            ["Recursive walk-forward", f"Mean RMSLE {walk['mean_rmsle']:.5f}", "reports/results/walk_forward_metrics.json"],
            ["Real Kaggle test submission", f"{submission['submission_rows']:,} rows", "submissions/final_submission.csv"],
        ],
        widths=[Inches(2.1), Inches(1.45), Inches(2.65)],
    )


def build_report() -> None:
    hpo = load_json(DATA / "hpo_final_comparison.json")
    quantile = load_json(DATA / "quantile_final_calibrated.json")
    walk = load_json(RESULTS / "walk_forward_metrics.json")
    submission = load_json(SUBMISSIONS / "final_submission_manifest.json")

    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "AI for Food Waste Reduction Model Refinement and Test Submission"
    doc.core_properties.subject = "Samsung Innovation Campus AI 17 Capstone Group 11 final report"
    doc.core_properties.author = "SIC AI 17 Capstone Group 11"
    doc.core_properties.keywords = "LightGBM, demand forecasting, food waste, model refinement, submission"

    add_title_page(doc, walk, submission)

    doc.add_page_break()
    doc.add_heading("Model Refinement", level=1)

    doc.add_heading("1 Overview", level=2)
    add_paragraph(
        doc,
        "Refinement çalışması dört katmandan oluşur: mevcut LightGBM modelinin yeniden üretilmesi, sızıntısız Optuna hiperparametre araması, quantile tahmin aralığının coverage kalibrasyonu ve gerçek submission koşullarını taklit eden recursive walk-forward doğrulaması. Tüm model seçimleri zaman sırasını korur. Hedef içermeyen Kaggle test kümesi hiçbir seçim metriğinde kullanılmaz."
    )
    add_paragraph(
        doc,
        "Nihai nokta tahmin modeli mevcut feature setini korur. Promotion etkileşimleri, uzun geçmiş özellikleri ve hybrid blend adayları ayrı deneylerde değerlendirildi. Önceden tanımlanan genel performans ve düşük talep guardrail'lerini birlikte geçemeyen adaylar final modele alınmadı."
    )

    doc.add_heading("2 Model Evaluation", level=2)
    add_paragraph(
        doc,
        "Baseline model aynı veri hazırlama ve zaman ayrımıyla yeniden çalıştırıldı. Sonuçlar önceki fazdaki değerlerle eşleşti. Optuna ile bulunan ayarlar holdout RMSLE'de %0.20, MAE'de %0.50 iyileşme sağladı. Kazanç sınırlıdır; ancak test dönemi hiperparametre seçimi sırasında kullanılmadığı için sonuç tarafsızdır."
    )
    add_table(
        doc,
        ["Metric", "Baseline", "Tuned", "Change"],
        [
            ["Holdout RMSLE", f"{hpo['baseline']['rmsle_test']:.5f}", f"{hpo['tuned']['rmsle_test']:.5f}", f"%{hpo['rmsle_improvement_pct']:.2f} better"],
            ["Holdout MAE", f"{hpo['baseline']['mae_test']:.2f}", f"{hpo['tuned']['mae_test']:.2f}", f"%{hpo['mae_improvement_pct']:.2f} better"],
            ["Best iteration", str(hpo['baseline']['best_iteration']), str(hpo['tuned']['best_iteration']), "Lower learning rate"],
        ],
        widths=[Inches(1.8), Inches(1.45), Inches(1.45), Inches(1.55)],
    )
    add_figure(
        doc,
        "refine_hpo_comparison.png",
        "Figure 1  Baseline and Optuna tuned holdout performance",
        alt_text="Bar charts comparing baseline and tuned LightGBM RMSLE and MAE",
    )

    doc.add_heading("3 Refinement Techniques", level=2)
    add_bullets(
        doc,
        [
            "Optuna TPE sampler ile 40 denemelik iç-validasyon araması yapıldı.",
            "P10-P90, P5-P95 ve P8-P92 quantile bantları coverage ve pinball loss ile karşılaştırıldı.",
            "Üç expanding-window fold, 10 haftalık recursive tahminle değerlendirildi.",
            "Promotion, uzun geçmiş ve hybrid adayları düşük talep performansını koruyan guardrail'lerle test edildi.",
        ],
    )
    add_paragraph(
        doc,
        "Uzun geçmiş adayının ortalama RMSLE'si %1.20 iyileşmesine rağmen son fold düşük talep RMSLE'si %1.99 kötüleşti. Hybrid aday genel RMSLE'yi %1.78 iyileştirdi; ancak düşük talep RMSLE'sindeki %0.65 gerileme %0.50 sınırını aştı. Bu nedenle iki aday da final submission'a alınmadı."
    )

    doc.add_heading("4 Hyperparameter Tuning", level=2)
    add_paragraph(
        doc,
        "Optuna araması haftalar 1-125 üzerinde eğitildi ve haftalar 126-135 üzerinde seçildi. Haftalar 136-145 arama boyunca kapalı tutuldu. Tuned model daha düşük learning rate ve daha fazla boosting iteration kullanarak daha kademeli bir öğrenme düzeni oluşturdu."
    )
    add_table(
        doc,
        ["Parameter", "Baseline", "Tuned"],
        [
            ["learning_rate", "0.05", "0.02692"],
            ["num_leaves", "63", "64"],
            ["min_data_in_leaf", "30", "48"],
            ["feature_fraction", "0.90", "0.70831"],
            ["bagging_fraction", "0.90", "0.84549"],
            ["bagging_freq", "1", "4"],
            ["lambda_l1", "0", "0.01920"],
            ["lambda_l2", "0", "0.000055"],
        ],
        widths=[Inches(2.4), Inches(1.8), Inches(1.8)],
    )
    add_code(
        doc,
        [
            "inner_train = data[data.week <= 125]",
            "inner_valid = data[(data.week >= 126) & (data.week <= 135)]",
            "study.optimize(objective, n_trials=40)",
        ],
    )

    doc.add_heading("5 Cross Validation", level=2)
    add_paragraph(
        doc,
        "Final doğrulama, eğitim bitiş haftaları 115, 125 ve 135 olan üç expanding-window fold kullanır. Her fold sonraki 10 haftayı tahmin eder. İlk doğrulama haftasının lag değerleri geçmiş gerçek talepten, sonraki haftaların lag değerleri ise önceki model tahminlerinden oluşturulur. Böylece statik 10 haftalık submission koşulu taklit edilir ve gelecekteki hedeflerden sızıntı önlenir."
    )
    walk_rows = []
    for fold in walk["folds"]:
        walk_rows.append(
            [
                f"{fold['validation_start_week']}-{fold['validation_end_week']}",
                f"{fold['rmsle']:.5f}",
                f"{fold['mae']:.2f}",
                f"{fold['persistence_rmsle']:.5f}",
                f"%{fold['rmsle_improvement_vs_persistence_pct']:.2f}",
            ]
        )
    walk_rows.append(
        [
            "Mean",
            f"{walk['mean_rmsle']:.5f}",
            f"{walk['mean_mae']:.2f}",
            f"{walk['mean_persistence_rmsle']:.5f}",
            f"%{walk['mean_rmsle_improvement_vs_persistence_pct']:.2f}",
        ]
    )
    add_table(
        doc,
        ["Validation weeks", "Model RMSLE", "Model MAE", "Persistence RMSLE", "Improvement"],
        walk_rows,
        widths=[Inches(1.3), Inches(1.2), Inches(1.05), Inches(1.45), Inches(1.2)],
    )
    add_figure(
        doc,
        "final_walk_forward_validation.png",
        "Figure 2  Recursive LightGBM and persistence baseline across three validation folds",
        alt_text="Grouped bar chart showing model and persistence RMSLE for three walk-forward folds",
    )
    add_paragraph(
        doc,
        "Recursive ortalama RMSLE 0.49833, tek holdout RMSLE 0.47534'ten yüksektir. Bu fark model bozulması anlamına gelmez; recursive protokol daha zordur çünkü tahmin hataları sonraki haftaların lag özelliklerine taşınır."
    )

    doc.add_heading("6 Feature Selection", level=2)
    add_paragraph(
        doc,
        "Final feature seti fiyat, promosyon, takvim, merkez, yemek ve geçmiş talep bilgilerini içeren 19 değişkenden oluşur. Mevcut feature importance sonuçlarında num_orders_lag_1 ve num_orders_roll_mean_4 baskındır. Promotion etkileşimleri genel sonucu iyileştirmedi. Uzun geçmiş özellikleri ortalama sonucu iyileştirse de düşük talep diliminde guardrail'i geçtiği için reddedildi. Bu karar, düşük talep ürünlerinde aşırı tahminin gıda israfı riskini artırabileceği iş hedefiyle uyumludur."
    )
    doc.add_page_break()
    doc.add_heading("Test Submission", level=1)

    doc.add_heading("1 Overview", level=2)
    add_paragraph(
        doc,
        "Test submission aşaması iki ayrı veri kümesini doğru biçimde ayırır. Haftalar 136-145 etiketli holdout dönemidir ve performans ölçümü için kullanılır. Kaggle test.csv ise haftalar 146-155'i içerir, hedef değişkeni yoktur ve yalnızca final submission üretiminde kullanılır. Final model tüm etiketli haftalar 1-145 üzerinde yeniden eğitilmiştir."
    )

    doc.add_heading("2 Data Preparation for Testing", level=2)
    add_paragraph(
        doc,
        "Train ve test tabloları meal_info ve fulfilment_center_info ile many-to-one olarak birleştirilir. Pipeline benzersiz kimlikleri, zorunlu kolonları, metadata eşleşmesini, negatif hedefleri ve center-meal-week tekrarlarını doğrular. Kategorik değerler train ve test üzerinde ortak, deterministik map ile kodlanır."
    )
    add_table(
        doc,
        ["Dataset", "Weeks", "Rows", "Purpose"],
        [
            ["Labeled training", "1-145", f"{submission['training_rows']:,}", "Final model training"],
            ["Holdout slice", "136-145", "32,821", "Reported model metrics"],
            ["Unlabeled Kaggle test", "146-155", f"{submission['submission_rows']:,}", "Final submission only"],
        ],
        widths=[Inches(1.7), Inches(1.2), Inches(1.25), Inches(2.1)],
    )
    add_paragraph(
        doc,
        "Lag özellikleri submission sırasında recursive üretilir. Hafta 146 için geçmiş etiketli talepler kullanılır. Hafta 147 ve sonrasında, hedefler bulunmadığından önceki haftanın model tahmini geçmişe eklenir."
    )

    doc.add_heading("3 Model Application", level=2)
    add_paragraph(
        doc,
        "Nihai LightGBM modeli tuned hiperparametrelerle ve 1.360 boosting round ile eğitildi. Tahminler log1p uzayından expm1 ile sipariş ölçeğine döndürüldü ve negatif değerler sıfırda sınırlandı. Submission, sample_submission kimlik sırasıyla one-to-one birleştirildi. Eksik, tekrarlı veya negatif sonuçlar pipeline tarafından reddedilir."
    )
    add_code(
        doc,
        [
            "model = train_point_model(train_observed, num_boost_round=1360)",
            "predictions = recursive_predict(model, test_static, train_static)",
            "submission = format_submission(sample_submission, predictions)",
        ],
    )

    doc.add_heading("4 Test Metrics", level=2)
    add_paragraph(
        doc,
        "Hedef içermeyen Kaggle test kümesi için yerel RMSLE veya MAE hesaplanamaz. Model kalitesi etiketli holdout ve recursive walk-forward sonuçlarıyla raporlanır. Test kümesinde yalnızca submission bütünlüğü ve tahmin dağılımı kontrol edilir."
    )
    add_table(
        doc,
        ["Submission check", "Result"],
        [
            ["Rows", f"{submission['submission_rows']:,}"],
            ["Week range", f"{submission['test_week_range'][0]}-{submission['test_week_range'][1]}"],
            ["Minimum prediction", f"{submission['prediction_min']:.2f}"],
            ["Mean prediction", f"{submission['prediction_mean']:.2f}"],
            ["Maximum prediction", f"{submission['prediction_max']:.2f}"],
            ["Missing predictions", "0"],
            ["Negative predictions", "0"],
        ],
        widths=[Inches(3.25), Inches(2.75)],
    )
    add_figure(
        doc,
        "final_submission_weekly_mean.png",
        "Figure 3  Weekly mean demand in the final submission",
        width=5.9,
        alt_text="Line chart of mean predicted orders for weeks 146 through 155",
    )

    doc.add_heading("5 Model Deployment", level=2)
    add_paragraph(
        doc,
        "Teslim kapsamında üretim API entegrasyonu yapılmamıştır. Bunun yerine yeniden kullanılabilir inference pipeline'ı, tam veriyle eğitilmiş model, final submission ve üretim manifesti sağlanmıştır. Manifest eğitim ve test dönemlerini, satır sayılarını, boosting round değerini ve haftalık tahmin özetlerini kaydeder. Bu dosyalar yeni veri geldiğinde haftalık yeniden eğitim veya dinamik indirim sistemine entegrasyon için temel oluşturur."
    )
    add_table(
        doc,
        ["Artifact", "Location", "Status"],
        [
            ["Final recursive model", "data/final_submission_lgb_model.txt", "Generated"],
            ["Final Kaggle submission", "submissions/final_submission.csv", "Validated"],
            ["Submission manifest", "submissions/final_submission_manifest.json", "Generated"],
            ["Walk-forward metrics", "reports/results/walk_forward_metrics.json", "Generated"],
        ],
        widths=[Inches(1.85), Inches(3.1), Inches(1.15)],
    )

    doc.add_heading("6 Code Implementation", level=2)
    add_paragraph(
        doc,
        "Kod iki katmanda düzenlenmiştir. notebooks klasöründeki betikler önceki fazları ve refinement deneylerini yeniden üretir. extensions/forecasting klasörü ise sızıntısız veri doğrulaması, recursive inference, walk-forward değerlendirme, guarded deneyler, hata analizi ve submission üretimini sağlar."
    )
    add_table(
        doc,
        ["Component", "Responsibility"],
        [
            ["notebooks/repro_01_02.py", "Data preparation and feature engineering"],
            ["notebooks/refine_01_hpo.py", "Leakage-free Optuna search"],
            ["extensions/forecasting/pipeline.py", "Validated preparation, training and recursive prediction"],
            ["extensions/forecasting/walk_forward_validation.py", "Three-fold expanding-window backtest"],
            ["extensions/forecasting/build_submission.py", "Full-data model and final submission"],
            ["extensions/forecasting/tests/test_pipeline.py", "Unit tests for leakage and submission invariants"],
        ],
        widths=[Inches(2.8), Inches(3.2)],
    )
    add_paragraph(
        doc,
        "Dokuz birim test lag zamanlamasını, recursive history güncellemesini, uzun geçmiş özelliklerini, submission sırasını, persistence baseline'ı, metrikleri, hata dilimlerini ve hybrid yardımcılarını denetler. Final teslim öncesinde dokuz testin tamamı başarıyla geçmiştir."
    )

    doc.add_heading("Conclusion", level=1)
    add_paragraph(
        doc,
        "Model refinement ve test submission aşamaları tamamlanmıştır. Optuna tuning küçük fakat tarafsız bir holdout kazancı üretmiştir. Quantile kalibrasyonu %80 coverage hedefini %80.59 ile karşılamıştır. Recursive walk-forward doğrulaması, modelin 10 haftalık statik tahmin koşulunda persistence baseline'dan belirgin biçimde daha iyi olduğunu göstermiştir. Final model tüm etiketli veriyle eğitilmiş ve 32.573 satırlık gerçek Kaggle submission dosyası oluşturulmuştur."
    )
    add_paragraph(
        doc,
        "Sonuçların iki sınırlılığı vardır. Quantile alpha değeri ayrı bir kalibrasyon setinde yeniden seçilmemiştir. Ayrıca Kaggle test hedefleri paylaşılmadığı için final submission performansı yerel olarak ölçülemez. Bu nedenle nihai performans iddiası holdout ve walk-forward sonuçlarıyla sınırlıdır."
    )

    doc.add_heading("References", level=1)
    add_bullets(
        doc,
        [
            "Kaggle. Food Demand Forecasting dataset by kannanaikkal.",
            "Ke, G. et al. 2017. LightGBM A Highly Efficient Gradient Boosting Decision Tree. NeurIPS.",
            "Akiba, T. et al. 2019. Optuna A Next generation Hyperparameter Optimization Framework. KDD.",
            "Koenker, R. 2005. Quantile Regression. Cambridge University Press.",
            "Abdullah, N., Shaikh, A. K., and Almusharraf, A. 2025. Towards a sustainable retail food chain. Journal of Posthumanism 5 6 571 to 588.",
            "Project repository https://github.com/edasaruhan/SIC_AI_17_Capstone_Group_11",
        ],
    )

    doc.save(OUTPUT)
    print(f"Saved report: {OUTPUT}")


if __name__ == "__main__":
    build_report()
